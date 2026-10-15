import csv
import io
import zipfile
from datetime import datetime
from pathlib import Path

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from ledger.models import JournalEntry

def entries_for(company):
    return JournalEntry.objects.filter(company=company, state="POSTED").prefetch_related("lines__account", "documents").order_by("date", "number")

def rows_for(company):
    rows = []
    for entry in entries_for(company):
        for line in entry.lines.all():
            rows.append([f"E-{entry.number:06d}", entry.date.isoformat(), entry.journal, entry.description,
                         line.account.code, line.account.name, line.description,
                         float(line.debit), float(line.credit), entry.state])
    return rows

def export_xlsx(company):
    book = Workbook(); sheet = book.active; sheet.title = "Journal"
    sheet.append(["Pièce", "Date", "Journal", "Libellé", "Compte", "Intitulé", "Détail", "Débit", "Crédit", "État"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="145C48")
    for row in rows_for(company): sheet.append(row)
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for col, width in zip("ABCDEFGHIJ", [14, 13, 12, 36, 12, 30, 26, 18, 18, 16]): sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2, min_col=8, max_col=9):
        for cell in row: cell.number_format = '#,##0.00'
    stream = io.BytesIO(); book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="journal-{company.pk}-{timezone.localdate()}.xlsx"'
    return response

def export_pdf(company):
    stream = io.BytesIO(); doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet(); story = [Paragraph(company.name, styles["Title"]), Paragraph("Journal comptable — écritures comptabilisées", styles["Heading2"]), Spacer(1, 5*mm)]
    data = [["Pièce", "Date", "Journal", "Libellé", "Compte", "Intitulé", "Débit", "Crédit"]]
    for row in rows_for(company): data.append([*row[:6], f"{row[7]:,.2f}", f"{row[8]:,.2f}"])
    table = Table(data, repeatRows=1, colWidths=[22*mm, 20*mm, 19*mm, 65*mm, 20*mm, 55*mm, 26*mm, 26*mm])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#145c48")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),7),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#dce4e0")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f3f6f4")])]))
    story.append(table); doc.build(story)
    response = HttpResponse(stream.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="journal-{company.pk}-{timezone.localdate()}.pdf"'
    return response

def export_zip(company):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        csv_buffer = io.StringIO(newline="")
        writer = csv.writer(csv_buffer, delimiter=";"); writer.writerow(["Piece", "Date", "Journal", "Libelle", "Compte", "Intitule", "Detail", "Debit", "Credit", "Etat"])
        writer.writerows(rows_for(company)); archive.writestr("journal.csv", csv_buffer.getvalue().encode("utf-8-sig"))
        for entry in entries_for(company):
            for doc in entry.documents.all():
                try:
                    with doc.file.open("rb") as source: archive.writestr(f"pieces/E-{entry.number:06d}/{Path(doc.original_name).name}", source.read())
                except OSError as exc:
                    raise RuntimeError(f"Pièce manquante pour E-{entry.number:06d} : {doc.original_name}") from exc
    response = HttpResponse(stream.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="comptapieces-{company.pk}-{timezone.localdate()}.zip"'
    return response
